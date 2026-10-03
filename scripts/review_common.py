"""Shared text and styles for the reviewer spreadsheets (exercise review, launch set)."""
import json
import os

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
seed = json.load(open(os.path.join(ROOT, 'supabase/seed/exercises.json')))['exercises']
repair = json.load(open(os.path.join(ROOT, 'supabase/seed/repair_tests.json')))['tests']
catalog = json.load(open(os.path.join(ROOT, 'supabase/seed/joint_movements.json')))
en = json.load(open(os.path.join(ROOT, 'src/i18n/locales/en.json')))

FONT = 'Arial'
HEADER_FILL = PatternFill('solid', fgColor='121212')
INPUT_FILL = PatternFill('solid', fgColor='FFF2CC')  # reviewer cells
THIN = Side(style='thin', color='DDD8D0')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

BANDS = {
    'kid': '9+ (kids)',
    'teen': '13+ (teens)',
    'young': '18+',
    'adult': '30+',
    'mid': '45+',
    'senior': '60+',
    'elder': '75+',
}
POSITIONS = {'standing': 'Standing', 'with_support': 'With support', 'seated_only': 'Seated'}
PARTS = {
    'warmup_general': 'Warm-up: easy cardio',
    'warmup_mobility': 'Warm-up: mobility',
    'main': 'Main work',
    'finisher_cardio': 'Finisher: cardio',
    'finisher_mobility': 'Finisher: mobility',
    'cooldown_walk': 'Cool-down: walk',
    'cooldown_stretch': 'Cool-down: stretch',
    'cooldown_breathing': 'Cool-down: breathing',
}
# Detailed equipment names (improvements v1, C), from the English locale.
EQUIPMENT = en['equipment']
RISK = {**en['safety']['painAreas'], **en['safety']['conditions']}
MUSCLE = en['muscles']
ANATOMY = en['muscleAnatomy']

COLUMNS = [
    ('ID (slug)', 26),
    ('Exercise', 30),
    ('Muscles (role, emphasis 0–1)', 46),
    ('Equipment', 20),
    ('Level (1–5)', 10),
    ('Minimum age', 13),
    ('Positions', 22),
    ('Contraindications', 34),
    ('Cues', 48),
    ('Used as', 26),
    ('Joint movements (range)', 40),
    ('Shorter range allowed', 26),
    ('Approve / Fix (Aprovado/Corrigir)', 18),
    ('Comment (Comentário)', 44),
]
DECISION_COL = 13
COMMENT_COL = 14
MP = en['movementPain']
# Moves the reviewer must decide on, shown with the contraindications (QA R8-06).
REVIEW_FLAGS = {
    slug: 'REVIEWER: clinician decides whether this loads the knee / pivots the knee'
    for slug in (
        'rp_band_seated_hip_in_turn',
        'rp_band_seated_hip_out_turn',
        'rx_iso_hip_internal_rotation',
        'rx_iso_hip_external_rotation',
    )
}
RANGES = {'full': 'full', 'partial': 'partial', 'isometric': 'hold, no movement'}


def movement_text(joint, movement):
    return f"{MP['joints'][joint]}: {MP['movements'][joint][movement]['name']}"


def joints_text(joints):
    return '\n'.join(f'{movement_text(j, m)} ({RANGES[r]})' for j, m, r in joints) or '—'


def muscles_text(entry):
    lines = []
    for key, role, emphasis in entry['muscles']:
        lines.append(f'{MUSCLE[key]} ({ANATOMY[key]}): {role}, {emphasis:.1f}')
    return '\n'.join(lines) or '— (breathing, no muscle mapping)'


def row_for(entry):
    return [
        entry['slug'],
        en['exercises'][entry['slug']]['name'],
        muscles_text(entry),
        ', '.join(EQUIPMENT[q] for q in entry['equipment']) or 'Bodyweight',
        entry['level'],
        BANDS[entry['minAgeBand']],
        ', '.join(POSITIONS[p] for p in entry['positions']),
        (', '.join(RISK[c] for c in entry['contraindications']) or 'None')
        + (f"\n{REVIEW_FLAGS[entry['slug']]}" if entry['slug'] in REVIEW_FLAGS else ''),
        en['exercises'][entry['slug']]['cues'],
        ', '.join(PARTS[p] for p in entry['parts'])
        + (' · recovery plan only' if entry.get('rehab') else ''),
        joints_text(entry.get('joints', [])),
        '\n'.join(movement_text(*k.split('.')) for k in entry.get('rangeLimit', [])) or '—',
        None,
        None,
    ]


def style_header(ws, row):
    for col, (title, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=row, column=col, value=title)
        cell.font = Font(name=FONT, bold=True, color='FFFFFF')
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical='center')
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(col)].width = width


def style_row(ws, row, reviewer_cells=True):
    for col in range(1, len(COLUMNS) + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = Font(name=FONT)
        cell.alignment = Alignment(wrap_text=True, vertical='top')
        cell.border = BORDER
        if reviewer_cells and col in (DECISION_COL, COMMENT_COL):
            cell.fill = INPUT_FILL
